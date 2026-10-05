#!/usr/bin/env python3
"""Build the updated Korean PAAD manuscript report from aggregate outputs only."""
import argparse, csv, json, math
from pathlib import Path
from datetime import date
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from PIL import Image


def read_tsv(root, name):
    with (root / "tables" / name).open(encoding="utf-8", newline="") as f:
        return list(csv.DictReader(f, delimiter="\t"))


def num(x):
    try: return float(x)
    except (TypeError, ValueError): return float("nan")


def pval(x):
    x = num(x)
    if not math.isfinite(x): return "NE"
    return "<0.001" if x < .001 else f"{x:.3f}"


def hr_text(r):
    return f"HR {num(r['HR']):.2f} (95% CI {num(r['lower95']):.2f}–{num(r['upper95']):.2f})"


def lookup(rows, **kwargs):
    return next(r for r in rows if all(r.get(k) == str(v) for k, v in kwargs.items()))


def build_report(root: Path, target: Path):
    prov = json.loads((root / "private" / "provenance.json").read_text(encoding="utf-8"))
    flow = {r["item"]: int(float(r["n"])) for r in read_tsv(root, "Input_Cohort_flow.tsv")}
    f2g = read_tsv(root, "Input_Figure2_global_tests.tsv")
    f2p = read_tsv(root, "Input_Figure2_pairwise_tests.tsv")
    pg = read_tsv(root, "Input_Preop_platinum_global_comparisons.tsv")
    pd = read_tsv(root, "Input_Preop_platinum_descriptive.tsv")
    pc = read_tsv(root, "Input_Preop_platinum_Cox_models.tsv")
    plr = read_tsv(root, "Input_Preop_platinum_OS_logrank.tsv")[0]
    rg = read_tsv(root, "Input_Recurrence_pattern_global_comparisons.tsv")
    rp = read_tsv(root, "Input_Recurrence_pattern_pairwise_comparisons.tsv")
    rc = read_tsv(root, "Input_Recurrence_pattern_Cox_models.tsv")
    rlr = read_tsv(root, "Input_Recurrence_pattern_post_OS_logrank.tsv")[0]
    rlp = read_tsv(root, "Input_Recurrence_pattern_post_OS_pairwise_logrank.tsv")
    main_surv = read_tsv(root, "Table2_Main_survival_models.tsv")
    core_cox = read_tsv(root, "Input_3_Cox_coefficients.tsv")
    core_ph = read_tsv(root, "Input_3_PH_tests.tsv")
    core_lr = read_tsv(root, "Input_4_Logrank.tsv")
    core_lrp = read_tsv(root, "Input_4_Logrank_pairs.tsv")
    tmb = read_tsv(root, "Input_7_Reported_TMB.tsv")[0]
    repair_sets = read_tsv(root, "Input_Preop_platinum_HRD_MMR_set_tests.tsv")
    distant_lr = read_tsv(root, "Input_Distant_only_pattern_post_OS_logrank.tsv")[0]
    distant_cox = read_tsv(root, "Input_Distant_only_pattern_post_OS_Cox_models.tsv")
    expanded_lr = read_tsv(root, "Input_Expanded_recurrence_post_OS_logrank.tsv")[0]
    expanded_cox = read_tsv(root, "Input_Expanded_recurrence_post_OS_Cox_models.tsv")
    signature_qc = read_tsv(root, "Input_Signature_sample_quality_summary.tsv")
    signature_pooled = read_tsv(root, "Input_Preop_platinum_pooled_signature_sensitivity.tsv")
    signature_individual = read_tsv(root, "Input_Preop_platinum_pooled_signature_by_SBS.tsv")
    signature_status = read_tsv(root, "Input_Signature_requested_metrics_status.tsv")
    signature_assignment = read_tsv(root, "Input_Signature_target_assignment_audit.tsv")
    signature_audit = read_tsv(root, "Input_MAF_signature_audit.tsv")

    def desc(variable, group, category="Median [IQR]"):
        return lookup(pd, variable=variable, group=group, category=category)["value"]

    age_no, age_yes = desc("Age", "No"), desc("Age", "Yes")
    lvi_no, lvi_yes = desc("LVI", "No", "Positive"), desc("LVI", "Yes", "Positive")
    tp = lookup(f2g, gene="TP53", characteristic="Differentiation")
    cd3 = lookup(f2g, gene="CDKN2A", characteristic="N_stage")
    cd2 = lookup(f2g, gene="CDKN2A", characteristic="N_category")
    kras_lr = lookup(core_lr, variable="KRAS_group", endpoint="OS")
    kras_ph = next(r for r in core_ph if r.get("exposure") == "KRAS_group" and r.get("family") == "Extended genomic" and r.get("term") == "GLOBAL")
    sig_audit = {r["item"]: r["value"] for r in signature_audit}
    kras_source_label = ("Main 시트의 KRAS_subtype_MAF"
                         if prov.get("kras_source") == "workbook_maf"
                         else prov.get("kras_source_column", prov.get("kras_source", "미확인")))
    hrd_set = lookup(repair_sets, gene_set="HRD", outcome="Any retained variant")
    mmr_set = lookup(repair_sets, gene_set="MMR", outcome="Any retained variant")
    liver_lung = lookup(distant_cox, group1="Liver only", group2="Lung only")
    lung_other = lookup(distant_cox, group1="Lung only", group2="Other distant")
    liver_other = lookup(distant_cox, group1="Liver only", group2="Other distant")
    expanded_ld_lung = lookup(expanded_cox, group1="Local + distant", group2="Lung only")
    expanded_liver_lung = lookup(expanded_cox, group1="Liver only", group2="Lung only")

    doc = Document()
    sec = doc.sections[0]
    sec.top_margin = Inches(.62); sec.bottom_margin = Inches(.62)
    sec.left_margin = Inches(.72); sec.right_margin = Inches(.72)
    styles = doc.styles
    for style_name in ["Normal", "Title", "Subtitle", "Heading 1", "Heading 2", "Heading 3"]:
        style = styles[style_name]
        style.font.name = "Pretendard"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Pretendard")
    styles["Normal"].font.size = Pt(9.5)
    styles["Normal"].paragraph_format.space_after = Pt(4)
    styles["Normal"].paragraph_format.line_spacing = 1.12
    styles["Title"].font.size = Pt(24); styles["Title"].font.bold = True
    styles["Heading 1"].font.size = Pt(16); styles["Heading 1"].font.color.rgb = RGBColor(0,0,0)
    styles["Heading 2"].font.size = Pt(12); styles["Heading 2"].font.color.rgb = RGBColor(25,65,85)

    footer = sec.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = footer.add_run("PAAD manuscript analysis · updated clinical workbook · ")
    run.font.name = "Pretendard"; run.font.size = Pt(8)
    fld = OxmlElement("w:fldSimple"); fld.set(qn("w:instr"), "PAGE"); footer._p.append(fld)

    def para(text="", bold_lead=None, style=None):
        p = doc.add_paragraph(style=style)
        if bold_lead and text.startswith(bold_lead):
            r = p.add_run(bold_lead); r.bold = True
            p.add_run(text[len(bold_lead):])
        else: p.add_run(text)
        return p

    def heading(text, level=1, page=False):
        h = doc.add_heading(text, level=level)
        if page: h.paragraph_format.page_break_before = True
        return h

    def image(name, width=7.0):
        path = root / "figures" / name
        with Image.open(path) as im:
            w, h = im.size
        width = min(width, 8.0 * w / h)
        p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(); r.add_picture(str(path), width=Inches(width))
        for node in r._r.iter(qn("wp:docPr")): node.set("descr", name.replace("_", " "))

    def inputs(names):
        # Keep provenance visible without allowing a short filename tail to
        # occupy an otherwise blank page after a dense interpretation block.
        text = "Inputs: " + "; ".join(f"{i}. {name}" for i, name in enumerate(names, 1))
        p = para(text)
        p.paragraph_format.space_before = Pt(2)
        p.paragraph_format.space_after = Pt(2)
        for r in p.runs:
            r.font.size = Pt(7.3)
            r.font.color.rgb = RGBColor(85,85,85)

    def stat_table(headers, rows, widths=None):
        t = doc.add_table(rows=1, cols=len(headers)); t.alignment = WD_TABLE_ALIGNMENT.CENTER
        t.autofit = False
        if widths is None: widths = [7.0/len(headers)]*len(headers)
        for j,h in enumerate(headers):
            t.rows[0].cells[j].text = h; t.rows[0].cells[j].width = Inches(widths[j])
        for row in rows:
            cells = t.add_row().cells
            for j, value in enumerate(row): cells[j].text = str(value); cells[j].width = Inches(widths[j])
        for i,row in enumerate(t.rows):
            row._tr.get_or_add_trPr().append(OxmlElement("w:cantSplit"))
            for j,cell in enumerate(row.cells):
                cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
                shade = OxmlElement("w:shd"); shade.set(qn("w:fill"), "E6E6E6" if i==0 else ("F5F5F5" if i%2==0 else "FFFFFF")); cell._tc.get_or_add_tcPr().append(shade)
                for p in cell.paragraphs:
                    p.paragraph_format.space_after = Pt(0)
                    for r in p.runs: r.font.name="Pretendard"; r.font.size=Pt(8.2); r.bold=(i==0)
        return t

    doc.add_paragraph("Updated PDAC clinicogenomic analysis", "Title")
    para("260927_v3_PDAC_ANALYSIS_with_MAF_annotations.xlsx 적용 결과보고서", style="Subtitle")
    para(f"작성일 {date.today():%Y-%m-%d}  |  분석 코호트 {flow['Analysis patients']:,}명  |  전체 사망 {flow['OS deaths']:,}건")
    para(f"목적: 업데이트된 임상정보와 MAF 변이를 연결하여 PDAC의 분자–병리 연관성, 선행 platinum 노출, 재발양상 및 생존을 하나의 논문 흐름으로 평가한다. KRAS 아형은 {kras_source_label}을 사용한다. 임상정보 TMB를 우선 사용하고 Variant count는 oncoplot의 보조 annotation으로 사용한다.")

    heading("핵심 결론", 1)
    para("1. 1,011명의 전체 병기 환자 코호트에서 TP53 변이 검출률은 분화도에 따라, CDKN2A 변이 검출률은 림프절 병기에 따라 달랐다. 이는 유전자 상태와 수술 병리 표현형이 연관될 수 있음을 보여주지만 인과관계를 뜻하지 않는다.")
    para(f"2. 선행 platinum Yes 301명과 No 706명의 OS log-rank p는 {pval(plr['p'])}였다. 전체 및 선행치료 환자 제한 Cox의 4개 사전지정 비교를 Holm 보정한 뒤에는 OS와 기록된 재발 모두 유의한 차이가 없었다. 따라서 이 자료만으로 platinum의 치료효과 또는 위해를 주장할 수 없다.")
    para(f"3. 재발 환자 659명에서 Local only 234명, Distant only 306명, Local + distant 119명이었다. 재발 후 생존의 전체 log-rank p는 {pval(rlr['p'])}였고, 보정 Cox에서 Local + distant 대비 Local only의 사망위험은 {hr_text(lookup(rc, group1='Local only', group2='Local + distant'))}, Distant only는 {hr_text(lookup(rc, group1='Distant only', group2='Local + distant'))}였다. 즉 복합 국소+원격 재발은 재발 후 불량한 임상경과를 표시하는 강한 예후표지자이다.")
    para("4. KRAS 아형은 전체 OS에서 이질성을 보였으며, 보정 Cox에서 G12D는 G12V 및 G12R보다 높은 사망 hazard와 연관되었다. 반면 주요 driver의 단순 변이 유무는 다중보정 후 독립적 차이가 제한적이었다. 아형 수준 분석이 이분형 검출/미검출보다 더 많은 예후 정보를 제공한다.")
    para(f"5. Distant only 306명 내부에서 Lung only는 Liver only 및 Other distant보다 재발 후 경과가 양호했다. 보정 HR은 Liver only/Lung only {hr_text(liver_lung)}, Lung only/Other distant {hr_text(lung_other)}였고 두 비교 모두 Holm 보정 후 유의했다. 반면 Liver only/Other distant {hr_text(liver_other)}는 차이를 지지하지 않았다.")
    para(f"6. Preoperative platinum Yes/No 사이에서 HRD-list 변이 검출률은 {num(hrd_set['yes_value']):.1f}%/{num(hrd_set['no_value']):.1f}%(Holm p={pval(hrd_set['p_holm'])}), MMR-list는 {num(mmr_set['yes_value']):.1f}%/{num(mmr_set['no_value']):.1f}%(Holm p={pval(mmr_set['p_holm'])})로 차이를 지지하지 않았다. 환자별 SBS signature 분석은 품질기준을 통과한 환자가 0명이어서 양성률·burden·proportion 비교를 산출하지 않았다.")

    heading("분석 대상과 통계 방법", 1, page=True)
    para(f"입력 파일은 data/raw/260927_v3_PDAC_ANALYSIS_with_MAF_annotations.xlsx의 Main 시트이다. 원본 {flow['Source records']:,}행에서 중복 patient_id에 해당하는 {flow['Duplicate-ID records excluded']}행(2개 ID)을 전부 제외하여 {flow['Analysis patients']:,}명을 분석하였다. M0 {flow['M0']}명과 M1 {flow['M1']}명을 포함하였다. 결과별 결측은 해당 분석에서만 제외하였다.")
    para("유전자 변이는 MAF의 유지된 variant call을 사용하였다. CNV가 필요한 LOH, tumor purity, tumor cellularity 분석은 수행하지 않았다. Oncoplot의 상단 Variant count는 환자별 전체 유지 변이 수이며, 유전자 오른쪽 count는 그 유전자의 유지 변이 수이다. 이는 TMB와 동일한 측정치가 아니다.")
    para("범주형 전체 비교는 Pearson χ² 또는 기대빈도가 작은 경우 Fisher exact를 사용하였다. 연속형 두 군 비교는 Wilcoxon rank-sum, 세 군 비교는 Kruskal–Wallis를 사용하였다. 세 군 이상의 모든 사전지정 쌍별 비교는 범주형 Fisher 또는 연속형 Dunn 검정을 사용하고 변수별 Holm 보정을 적용하였다. 여러 임상·유전체 변수의 전체 탐색에는 별도의 Holm 보정을 적용하였다.")
    para("생존곡선은 Kaplan–Meier, 전체 비교는 log-rank, 쌍별 비교는 Holm 보정 log-rank를 사용하였다. Cox 모형은 비교군/기준군의 HR과 95% CI를 보고하였다. HR>1은 앞에 적힌 비교군의 사건 순간위험이 기준군보다 높다는 뜻이다. KM log-rank는 비보정 곡선 비교이고 Cox는 공변량 보정 효과이므로 p값이 같을 필요가 없다. 전체 검정과 쌍별 검정, 분석 표본, 결측 제외 범위가 달라도 차이가 생길 수 있다.")
    para("보정 Cox는 연령 자연스플라인, 성별, T/N/M category, 선행치료, LVI, PNI 및 절제연을 포함하고 분화도별 기저위험을 허용하였다. 비례위험 가정은 Schoenfeld residual 기반 검정으로 확인하였다. 다중검정 q 또는 adjusted p와 95% CI가 서로 다른 범위를 보정하므로, adjusted p를 주된 유의성 판단으로 사용하였다.")
    para("원격재발 확장 분석은 Recurrence_pattern='Distant only' 환자에서 Distant_pattern을 Liver only, Lung only, Other distant로 재분류하였다. Peritoneum only, Multiple distant 및 그 밖의 원격재발은 모두 Other distant에 포함했다. 별도의 5군 분석은 Local only와 Local + distant를 유지하고 기존 Distant only를 위 세 군으로 분해하였다. KM은 전체 기록 환자, 보정 Cox는 공변량 완전사례를 사용하므로 표본수가 다를 수 있다.")
    para(f"SBS96은 {sig_audit['MAF file']}의 GRCh37 SNV만 사용해 SigProfilerMatrixGenerator 1.4.0과 COSMIC v3.6 exome reference로 분석했다. 전체 MAF에는 synonymous/silent 변이가 {sig_audit['Synonymous or silent rows']}개이고 panel BED/callable territory가 없어, 이 분석은 targeted-panel 탐색 분석으로 제한했다. 환자별 적합 기준은 context-counted SNV≥10, reconstruction cosine≥0.90, process count≥5, proportion≥0.20, bootstrap stability≥0.80이었다. 기준을 임의로 낮추지 않았으며, 집단 합산 refit은 환자 bootstrap CI와 label permutation p를 사용한 민감도 분석으로만 제시했다.")

    heading("통계값을 읽는 방법", 2, page=True)
    stat_table(
        ["통계값", "무엇을 나타내는가", "이 보고서에서의 해석"],
        [
            ["p-value", "차이 또는 연관성이 없다는 가정과 관측자료가 얼마나 부합하는지 평가한다.", "작을수록 귀무가설과 덜 부합한다. 효과크기나 임상적 중요도를 나타내지는 않는다."],
            ["Adjusted p", "여러 검정을 동시에 시행해 우연히 작은 p가 나올 가능성을 보정한다.", "Holm은 family-wise error, BH는 false discovery rate를 통제한다. 다중비교가 있으면 raw p보다 adjusted p를 우선한다."],
            ["95% CI", "자료와 모형이 허용하는 효과크기의 불확실성 범위이다.", "HR의 CI가 1을 포함하면 증가와 감소 양쪽 가능성이 남아 차이를 확정하기 어렵다. 좁을수록 추정이 정밀하다."],
            ["Cramer's V", "Cramér는 이 통계량을 정립한 Harald Cramér의 이름이고, V는 두 범주형 변수 사이 연관성의 크기를 나타내는 효과크기이다.", "0은 연관 없음, 1은 완전한 연관이다. 관례적으로 약 0.10은 small, 0.30은 moderate, 0.50은 large로 읽지만 표의 크기와 임상 맥락에 따라 달라지는 참고 기준이다."],
            ["Hazard ratio", "비교군의 순간 사건 hazard를 기준군과 비교한 비율이다.", "1은 차이 없음, 1보다 크면 비교군 hazard 증가, 작으면 감소이다. 누적 사망확률이나 평균 생존기간의 비율은 아니다."],
            ["Log-rank p", "보정하지 않은 생존곡선 전체가 같은지 검정한다.", "낮으면 하나 이상의 곡선이 다르다는 뜻이다. 어느 쌍이 다른지는 쌍별 검정으로 확인해야 한다."],
            ["PH test p", "Cox의 HR이 추적기간 동안 일정하다는 proportional hazards 가정을 점검한다.", "0.05 미만이면 가정 위반 가능성을 시사한다. 높은 p는 위반 근거가 없다는 뜻이지 가정을 입증하는 것은 아니다."],
        ],
        widths=[1.25, 2.75, 3.0],
    )

    heading("Table 1. 코호트 특성", 1, page=True)
    image("Table1_Cohort_characteristics.png", 7.15)
    para("표 해석: TMB는 836명에서 관측되고 175명에서 결측이다. 단위·검사법·버전·결측 원인을 확인할 수 없어 연속형 탐색 변수로만 사용하였다. Preoperative platinum의 Unknown 4명은 해당 비교에서 제외하였다.")
    inputs(["Table1_Cohort_characteristics.tsv", "Input_1_Missingness.tsv"])

    heading("Figure 1. 연구 코호트와 분자적 지형", 1, page=True)
    for label, title, fn in [
        ("1a", "환자 선정 흐름", "Figure1A_Cohort.png"),
        ("1b", "MAF 변이와 임상정보 oncoplot", "Figure1B_Oncoplot.png"),
        ("1c", "KRAS 아형 분포", "Figure1C_KRAS_distribution.png")]:
        heading(f"{label}. {title}", 2); image(fn, 7.15)
    para(f"분석 방법: 각 열은 중복 환자 제외 후 한 환자이다. KRAS 유전자 행을 포함한 유전자 행은 MAF의 실제 변이유형을 표시하며, KRAS subtype annotation과 관련 분석은 {kras_source_label}을 사용한다. 환자별 총 Variant count와 각 유전자의 variant count를 동시에 표시하였다.")
    para("해석: 대규모 수술 중심 PDAC 코호트에서 주요 driver의 높은 검출 빈도와 임상적 이질성을 한 화면에 연결한다. Variant count는 기술적·탐색적 mutation burden이며 임상정보 TMB를 대체하지 않는다. 유전자별 panel coverage를 확인하지 못했으므로 미검출을 확정적인 생물학적 wild type으로 표현하지 않는다.")
    inputs(["Input_Cohort_flow.tsv", "Input_Figure1B_Oncoplot_gene_counts.tsv", "Input_Figure1B_Oncoplot_clinical_tracks.tsv", "Input_KRAS_distribution.tsv"])

    heading("Figure 2. 주요 driver와 병리 표현형", 1, page=True)
    panels2 = [("2a","유전자–병리 전체 연관성","Figure2A_Pathology_matrix.png"),
      ("2b","TP53 변이 검출과 분화도","Figure2_TP53_differentiation.png"),
      ("2c","CDKN2A 변이 검출과 N stage 0/1/2","Figure2_CDKN2A_N_stage.png"),
      ("2d","CDKN2A 변이 검출과 N category 0/1","Figure2_CDKN2A_N_category.png")]
    for lab,title,fn in panels2: heading(f"{lab}. {title}",2); image(fn,6.7)
    para(f"사용 컬럼과 통계: TP53_variant/MAF TP53 status와 Differentiation, CDKN2A_variant/MAF CDKN2A status와 N_stage 및 N_category를 비교하였다. TP53–분화도의 전체 χ² p={pval(tp['p'])}, Cramer's V={num(tp['cramers_v']):.2f}; CDKN2A–N stage p={pval(cd3['p'])}, V={num(cd3['cramers_v']):.2f}; CDKN2A–N category p={pval(cd2['p'])}, V={num(cd2['cramers_v']):.2f}였다. 3범주 쌍별 Fisher 검정은 각 변수 내 Holm 보정하였다. Cramér는 통계학자 Harald Cramér의 이름이고 V는 범주형 변수 사이 연관성의 크기를 0에서 1로 나타내는 효과크기이다. 0은 연관 없음, 1은 완전한 연관이다.")
    para("비교 방향: 막대의 분모는 각 분화도 또는 N 범주의 환자 수이고, 청색 부분은 Variant detected이다. 따라서 Figure 2b는 WD/MD/PD 사이에서 TP53 검출률이 다른지, Figure 2c–d는 N category 사이에서 CDKN2A 검출률이 다른지를 직접 보여준다. Variant detected군 내부의 병리 분포를 비교하는 그림이 아니다.")
    para(f"효과크기 해석: 관례적인 참고 기준은 V 약 0.10 small, 0.30 moderate, 0.50 large이다. 따라서 TP53–분화도의 V={num(tp['cramers_v']):.2f}, CDKN2A–N stage의 V={num(cd3['cramers_v']):.2f}, CDKN2A–N category의 V={num(cd2['cramers_v']):.2f}는 모두 small association에 해당한다. p-value는 낮지만 연관성의 크기는 약하다. 이는 1,011명의 큰 표본에서 작은 차이도 통계적으로 명확해질 수 있음을 보여준다.")
    para("결과 해석: TP53 검출률은 WD보다 MD와 PD에서 높았고, CDKN2A 검출률은 N0보다 N1에서 높았다. 통계적으로 차이는 확인되지만 효과크기가 작으므로 강한 병리 결정인자로 표현해서는 안 된다. 수술 후 병리와의 연관성이므로 진단 시점의 독립 예측이나 기전도 입증하지 않는다.")
    inputs(["Input_2_Pathology_CMH.tsv", "Input_Figure2_global_tests.tsv", "Input_Figure2_pairwise_tests.tsv", "Input_Figure2_counts.tsv"])

    heading("Figure 3. 선행 platinum 노출의 임상적 맥락과 결과", 1, page=True)
    for lab,title,fn in [("3a","임상·유전체 특성 비교","Figure3_Platinum_context.png"),
      ("3b","전체생존 KM","Figure3_Platinum_OS.png"),
      ("3c","OS와 기록된 재발 Cox","Figure3_Platinum_forest.png")]:
        heading(f"{lab}. {title}",2); image(fn,6.9)
    age = lookup(pg, variable="Age"); lvi = lookup(pg, variable="LVI")
    para(f"사용 컬럼과 모집단: Preop_platinum_exposure의 Yes 301명과 No 706명을 비교하고 Unknown 4명은 제외하였다. Yes는 모두 Neoadjuvant='y'였고 No에는 선행치료 미시행 664명과 비-platinum 선행치료 42명이 포함된다. 따라서 전체 비교는 치료 선택 맥락을 포함하는 기술적 비교이고, 보정 Cox는 선행치료 환자 342명에서 민감도 분석하였다.")
    para(f"기초 특성: 연령은 No {age_no}, Yes {age_yes}였고 Wilcoxon p={pval(age['p'])}, 15개 변수 Holm p={pval(age['p_holm'])}였다. LVI 양성은 No {lvi_no}, Yes {lvi_yes}였고 χ² p={pval(lvi['p'])}, Holm p={pval(lvi['p_holm'])}였다. 다른 유전체 지표(TMB, Variant count, 주요 driver)는 보정 후 유의한 차이가 없었다.")
    para(f"생존곡선 해석: 전체 OS log-rank p={pval(plr['p'])}는 보정하지 않은 Yes와 No 생존곡선의 전체 차이가 통상 기준 0.05에 미치지 못했음을 뜻한다. '차이가 전혀 없다'는 증명은 아니며, 관측된 차이를 우연 변동과 구분할 근거가 충분하지 않다는 의미이다.")
    para(f"HR 해석: HR은 Yes/No이다. 전체 비보정 OS {hr_text(pc[0])}는 Yes군의 순간 사망 hazard가 17% 높다는 추정이지만 95% CI가 1을 포함한다. 선행치료 환자 제한 보정 OS {hr_text(pc[1])}는 Yes군의 hazard가 23% 낮다는 추정이지만 CI가 1을 포함한다. 기록된 재발은 전체 {hr_text(pc[2])}로 16% 증가 방향, 제한 보정 {hr_text(pc[3])}로 33% 감소 방향이지만 두 CI 모두 1을 포함한다. 4개 모형의 Holm 보정 p가 모두 {pval(pc[0]['p_holm_all_models'])}이므로 어느 방향도 통계적으로 확정되지 않았다.")
    para(f"비례위험 가정: 전체 환자의 비보정 재발 모형은 PH global p={pval(pc[2]['PH_global_p'])}로 0.05보다 작아 HR이 시간에 따라 일정하다는 가정에 위반 근거가 있었다. 따라서 이 모형의 HR 1.16을 전체 추적기간에 동일한 효과로 해석하면 안 된다. 선행치료 환자 제한 보정 재발 모형은 PH global p={pval(pc[3]['PH_global_p'])}로 위반 근거가 없었지만, 높은 p가 가정의 참을 입증하는 것은 아니다.")
    para("임상 해석: platinum 노출군의 환자선택과 치료 맥락이 뚜렷하여 전체 KM은 치료효과 추정이 아니다. 선행치료 환자 내부 분석에서도 신뢰구간이 넓고 다중보정 후 차이가 없었으며, 전체 재발 모형에는 비례위험 위반도 있어 이 자료로 platinum이 생존 또는 재발을 개선·악화한다고 결론내릴 수 없다. 향후 regimen, dose intensity, response, 수술 전후 시간과 적응증 정보가 필요하다.")
    heading("3d. 선행 platinum에 따른 HRD/MMR gene-list 변이", 2); image("Figure3_Platinum_HRD_MMR.png", 6.0)
    para(f"분석 방법: McIntyre 논문의 사전지정 HRD 18개, MMR 4개 gene list에서 MAF retained variant가 하나 이상 검출된 비율을 Preop_platinum_exposure Yes/No로 비교했다. Gene-list 양성률은 Fisher exact, 변이 유전자 수는 Wilcoxon rank-sum을 사용하고 4개 set-level 검정을 Holm 보정했다. 개별 유전자는 gene list별 BH 보정했다.")
    para(f"결과: HRD-list 검출률은 No {num(hrd_set['no_value']):.1f}%, Yes {num(hrd_set['yes_value']):.1f}%, OR Yes/No {num(hrd_set['effect']):.2f}, Fisher p={pval(hrd_set['p'])}, Holm p={pval(hrd_set['p_holm'])}였다. MMR-list는 No {num(mmr_set['no_value']):.1f}%, Yes {num(mmr_set['yes_value']):.1f}%, OR {num(mmr_set['effect']):.2f}, p={pval(mmr_set['p'])}, Holm p={pval(mmr_set['p_holm'])}였다. 개별 유전자도 BH 보정 후 유의한 차이가 없었다.")
    heading("3d 해석 및 한계", 2, page=True)
    para(f"효과크기 해석: HRD-list의 OR {num(hrd_set['effect']):.2f}는 Yes군의 retained-variant 검출 odds가 No군과 거의 같음을 뜻한다. MMR-list의 OR {num(mmr_set['effect']):.2f}도 1에 가깝다. 두 비교의 원 p와 Holm 보정 p가 모두 0.05보다 크므로, 관찰된 작은 차이를 표본 변동과 구분할 근거가 없다. 이는 두 군이 완전히 동일하다는 증명은 아니다.")
    para("양성 정의의 한계: 여기서 HRD/MMR 양성은 해당 gene list에 retained variant가 하나 이상 있다는 기술적 정의이다. 변이 등급, 양쪽 대립유전자 소실, germline 여부, MSI 또는 기능적 repair deficiency를 통합하지 않았으므로 임상 바이오마커 양성률로 해석하면 안 된다.")
    para("해석: platinum 노출군과 비노출군의 baseline retained-variant 분포가 HRD/MMR gene-list 수준에서 뚜렷하게 다르다는 근거는 없다. 그러나 이 결과는 변이의 병원성, biallelic loss, LOH, germline, MSI 또는 기능적 HRD/dMMR를 평가한 것이 아니다. 또한 치료 선택이 무작위가 아니므로 platinum 감수성 또는 치료효과를 뜻하지 않는다.")
    inputs(["Input_Preop_platinum_population.tsv", "Input_Preop_platinum_descriptive.tsv", "Input_Preop_platinum_global_comparisons.tsv", "Input_Preop_platinum_pairwise_comparisons.tsv", "Input_Preop_platinum_Cox_models.tsv", "Input_Preop_platinum_OS_logrank.tsv", "Input_Preop_platinum_HRD_MMR_set_tests.tsv", "Input_Preop_platinum_HRD_MMR_gene_tests.tsv"])

    heading("Figure 4. 재발양상과 재발 후 임상경과", 1, page=True)
    for lab,title,fn in [("4a","재발양상 구성","Figure4_Recurrence_distribution.png"),
      ("4b","재발양상별 임상·유전체 비교","Figure4_Recurrence_context.png"),
      ("4c","재발 후 생존 KM","Figure4_Post_recurrence_OS.png"),
      ("4d","보정 재발 후 생존 Cox","Figure4_Post_recurrence_forest.png")]:
        heading(f"{lab}. {title}",2); image(fn,6.9)
    age_r = lookup(rg, variable="Age")
    age_pair = lookup(rp, variable="Age", group1="Distant only", group2="Local + distant")
    para(f"사용 컬럼과 모집단: Recurrence_event=1이고 Recurrence_pattern이 기록된 659명만 포함하였다. Distant only 306명(46.4%), Local only 234명(35.5%), Local + distant 119명(18.1%)이었다. 비재발 환자의 공란은 정상적인 비해당 값이므로 네 번째 범주로 만들지 않았다.")
    para(f"세 군 비교: 14개 임상·유전체 변수의 전체 검정에서 연령의 Kruskal–Wallis 원 p={pval(age_r['p'])}였으나 변수군 Holm p={pval(age_r['p_holm'])}였다. 연령의 사전지정 Dunn 쌍별 비교 중 Distant only 대 Local + distant는 변수 내 Holm p={pval(age_pair['p_holm'])}였지만, 상위 전체 검정의 광범위 보정을 통과하지 못했으므로 탐색적 결과로 해석한다. TMB와 Variant count를 포함한 유전체 변수는 재발양상별 차이를 지지하지 않았다.")
    local_hr = lookup(rc, group1='Local only', group2='Local + distant')
    distant_hr = lookup(rc, group1='Distant only', group2='Local + distant')
    para(f"재발 후 생존의 전체 검정: OS_months−DFS_months를 재발 후 시간으로 정의하였다. 전체 log-rank p={pval(rlr['p'])}는 세 생존곡선 중 하나 이상이 다르다는 뜻이다. 이 값만으로 세 쌍이 모두 다르다고 결론내리지는 않으므로 쌍별 log-rank를 확인했으며, 모든 쌍에서 Holm p<0.05였다.")
    para(f"보정 HR 해석: Cox 전체 LRT p={pval(rc[0]['global_LRT_p'])}, PH global p={pval(rc[0]['PH_global_p'])}였다. Local only/Local + distant는 {hr_text(local_hr)}로 Local only의 순간 사망 hazard가 42% 낮았다. 방향을 뒤집으면 Local + distant의 hazard는 Local only의 약 {1/num(local_hr['HR']):.2f}배이다. Distant only/Local + distant는 {hr_text(distant_hr)}로 Distant only의 hazard가 33% 낮았고, 반대로 Local + distant는 Distant only의 약 {1/num(distant_hr['HR']):.2f}배였다. 두 CI는 모두 1을 포함하지 않아 Local + distant의 불량한 재발 후 경과를 지지한다.")
    para("임상 해석: 복합 국소+원격 재발은 단일 부위 재발보다 재발 후 경과가 불량한 고위험 임상표현형이다. 이 결과는 진단 시점 예측자가 아니라 재발 후에 정의되는 예후표지자이며, 치료선택·재발검사 빈도·구제치료의 영향을 받을 수 있다. 따라서 재발 예방 효과나 특정 치료의 인과효과로 해석하지 않는다.")
    inputs(["Input_Recurrence_pattern_distribution.tsv", "Input_Recurrence_pattern_descriptive.tsv", "Input_Recurrence_pattern_global_comparisons.tsv", "Input_Recurrence_pattern_pairwise_comparisons.tsv", "Input_Recurrence_pattern_post_OS_logrank.tsv", "Input_Recurrence_pattern_post_OS_pairwise_logrank.tsv", "Input_Recurrence_pattern_Cox_models.tsv"])

    heading("Figure 4 확장. 원격재발 부위와 재발 후 경과", 1, page=True)
    for lab,title,fn in [("4e","Distant only 내부 KM","Figure4_Distant_only_pattern_KM.png"),
      ("4f","Distant only 내부 보정 Cox","Figure4_Distant_only_pattern_forest.png"),
      ("4g","5군 확장 재발표현형 KM","Figure4_Expanded_recurrence_KM.png"),
      ("4h","5군 확장 재발표현형 보정 Cox","Figure4_Expanded_recurrence_forest.png")]:
        heading(f"{lab}. {title}",2); image(fn,6.9)
    para(f"Distant only 내부 분석: Liver only 133명, Lung only 60명, Other distant 113명이었다. 전체 log-rank p={pval(distant_lr['p'])}로 하나 이상의 곡선이 달랐다. 보정 Cox는 완전사례 {int(num(distant_cox[0]['n']))}명을 사용했고 global LRT p={pval(distant_cox[0]['global_LRT_p'])}, PH global p={pval(distant_cox[0]['PH_global_p'])}였다.")
    para(f"비교 방향: Liver only/Lung only {hr_text(liver_lung)}, Holm p={pval(liver_lung['p_holm'])}로 Liver only의 재발 후 사망 hazard가 Lung only보다 약 {num(liver_lung['HR']):.2f}배 높았다. Lung only/Other distant {hr_text(lung_other)}, Holm p={pval(lung_other['p_holm'])}로 Lung only의 hazard가 {100*(1-num(lung_other['HR'])):.1f}% 낮았다. Liver only/Other distant {hr_text(liver_other)}, Holm p={pval(liver_other['p_holm'])}로 차이를 지지하지 않았다. 즉 이 자료에서는 Lung only가 상대적으로 양호하고 Liver only와 Other distant는 비슷한 경과를 보였다.")
    para(f"5군 확장 분석: 전체 log-rank p={pval(expanded_lr['p'])}였다. 10개 Cox 대비를 Holm 보정한 결과 Local + distant/Lung only {hr_text(expanded_ld_lung)}, Holm p={pval(expanded_ld_lung['p_holm'])}, Liver only/Lung only {hr_text(expanded_liver_lung)}, Holm p={pval(expanded_liver_lung['p_holm'])}였다. 다만 Local + distant와 Liver only 또는 Other distant의 직접 대비는 보정 후 유의하지 않아, 'Local + distant가 모든 원격재발보다 항상 불량하다'고 확장해서는 안 된다.")
    para("임상적 의미: 원격재발을 하나의 Distant only로 묶으면 Lung-only의 상대적으로 양호한 자연경과가 희석된다. 따라서 향후 재발 후 예후모형과 임상시험 층화에서는 원격재발 부위를 분리할 근거가 있다. 그러나 재발 부위는 post-baseline 정보이며, 전이부담·구제치료·검사빈도와 같은 미측정 교란이 남아 진단 시점의 인과적 예측자로 사용할 수 없다.")
    inputs(["Input_Distant_pattern_grouping_rule.tsv", "Input_Distant_only_pattern_post_OS_distribution.tsv", "Input_Distant_only_pattern_post_OS_logrank.tsv", "Input_Distant_only_pattern_post_OS_pairwise_logrank.tsv", "Input_Distant_only_pattern_post_OS_Cox_models.tsv", "Input_Expanded_recurrence_post_OS_distribution.tsv", "Input_Expanded_recurrence_post_OS_logrank.tsv", "Input_Expanded_recurrence_post_OS_pairwise_logrank.tsv", "Input_Expanded_recurrence_post_OS_Cox_models.tsv"])

    heading("Figure 5. 주요 유전자 및 KRAS 아형의 예후 연관성", 1, page=True)
    for lab,title,fn in [("5a","주요 driver의 보정 OS","Figure3A_Gene_OS.png"),
      ("5b","KRAS 아형별 관찰 OS","Figure4A_KRAS_OS.png"),
      ("5c","KRAS 아형의 보정 쌍별 HR","Figure4B_KRAS_pairwise.png")]:
        heading(f"{lab}. {title}",2); image(fn,7.0)
    kd_v = next(r for r in main_surv if r['variable']=='KRAS subtype' and r['comparison_group']=='G12D' and r['reference_group']=='G12V')
    kd_r = next(r for r in main_surv if r['variable']=='KRAS subtype' and r['comparison_group']=='G12D' and r['reference_group']=='G12R')
    nd_d = next(r for r in main_surv if r['variable']=='KRAS subtype' and r['comparison_group']=='Not detected' and r['reference_group']=='G12D')
    para(f"전체 검정: KRAS 아형의 KM log-rank p={pval(kras_lr['p'])}는 다섯 아형의 보정하지 않은 생존곡선 중 하나 이상이 다르다는 뜻이다. 모든 아형 쌍이 서로 다르다는 의미는 아니므로 10개 쌍별 비교를 별도로 시행하고 Holm 보정하였다.")
    para(f"쌍별 HR 해석: HR은 앞의 비교군/뒤의 기준군이다. G12D/G12V는 {hr_text(kd_v)}, Holm p={pval(kd_v['adjusted_p'])}로 G12D의 순간 사망 hazard가 G12V보다 39% 높았다. G12D/G12R은 {hr_text(kd_r)}, Holm p={pval(kd_r['adjusted_p'])}로 G12D의 hazard가 G12R보다 67% 높았다. 두 CI가 모두 1을 포함하지 않고 다중검정 보정 후 p도 0.05보다 작아 해당 방향의 차이를 지지한다. Not detected/G12D는 {hr_text(nd_d)}, Holm p={pval(nd_d['adjusted_p'])}였다. HR은 누적 사망확률 또는 평균 생존기간의 비율이 아니라 특정 시점의 순간 사망 hazard 비율이다.")
    para(f"KM와 forest p값이 다른 이유: KM은 비보정 전체 또는 쌍별 log-rank이고 forest는 연령·병기·병리·치료 변수를 보정한 Cox 대비이다. 또한 전체 5군 가설과 10개 쌍별 가설, Holm 보정 범위가 다르다. 본 Cox 모형의 KRAS 항 전체 PH 검정 p={pval(kras_ph['p'])}였다. 따라서 두 결과는 모순이 아니라 서로 다른 질문에 답한다.")
    para("임상 해석: 단순 KRAS 검출/미검출보다 KRAS 아형의 이질성이 예후를 더 잘 구분했다. 특히 G12D의 불량 예후 연관성은 임상병리 보정 후에도 유지되었다. 다만 치료반응 예측, 표적치료 선택 또는 유전적 기전은 본 관찰자료만으로 입증할 수 없다.")
    inputs(["Input_3_Cox_coefficients.tsv", "Input_3_Model_diagnostics.tsv", "Input_3_PH_tests.tsv", "Input_4_Logrank.tsv", "Input_4_Logrank_pairs.tsv", "Input_4_Cox_pairs.tsv"])

    heading("Table 2. 주요 생존모형과 쌍별 대비", 1, page=True)
    image("Table2_Main_survival_models.png", 7.2)
    para("표의 HR은 Comparison / Reference이다. Raw p와 adjusted p를 함께 제시하며, 유전자 계수는 BH, KRAS 아형·재발양상 쌍은 Holm, platinum 4개 사전지정 모형은 Holm으로 보정하였다. Forest plot과 표는 같은 결과 객체에서 생성되어 수치가 일치한다.")
    inputs(["Table2_Main_survival_models.tsv"])

    heading("Supplementary figures", 1, page=True)
    supp = [("Supplementary Figure 1", "HRD 관련 유전자 변이 oncoplot", "Supplementary_Figure1_HRD.png",
             "HRD 유전자에서 관찰된 MAF variant와 임상 annotation을 기술한다. 기능적 HRD 진단 또는 LOH 결과가 아니다."),
            ("Supplementary Figure 2", "MMR 관련 유전자 변이 oncoplot", "Supplementary_Figure2_MMR.png",
             "MMR 유전자 variant 분포를 기술한다. MSI/dMMR 기능진단이나 germline 결과로 해석하지 않는다."),
            ("Supplementary Figure 3", "Reported TMB 탐색 분석", "Supplementary_Figure4_Reported_TMB.png",
             f"Reported TMB는 {int(float(tmb['observed']))}명에서 관측되고 {int(float(tmb['missing']))}명에서 결측이다. 단위와 검사법이 확인되지 않아 외부 cutoff를 적용하지 않았다.")]
    for sid,title,fn,note in supp:
        heading(f"{sid}. {title}",2); image(fn,7.0); para(note)
    heading("Supplementary Figure 4. SBS signature 분석 가능성 점검", 2)
    image("Supplementary_Figure_Signature_QC.png", 7.0)
    eligible_n = int(float(sig_audit.get("Cohort samples with >= 10 context-counted SNVs", 0)))
    para(f"MAF 점검 결과 전체 {flow['Analysis patients']}명 중 GRCh37 SBS96 context가 10개 이상인 환자는 {eligible_n}명이었다. 이 환자들의 COSMIC reconstruction cosine 최고값은 0.899로 사전지정 0.90 기준을 통과한 환자가 없었다. 따라서 환자별 SBS31/35, SBS3, MMR SBS의 positive rate, burden 및 proportion은 추정 불가(NE)로 처리하였다. 이는 'signature가 없다'는 결과가 아니라 현재 targeted-panel MAF로 신뢰성 있게 분류할 수 없다는 품질 결과이다.")
    para("원인과 의미: MAF에는 synonymous 변이가 없고 환자당 context-counted SNV 중앙값이 7개이며, panel BED와 callable territory도 없다. 이러한 선택된 저변이 입력에서는 COSMIC signature 분해가 불안정하고 panel trinucleotide opportunity 편향을 충분히 교정할 수 없다. 기준을 낮춰 양성률을 만들어내는 대신 미추정으로 남기는 것이 통계·생물정보학적으로 타당하다.")
    inputs(["Input_MAF_signature_audit.tsv", "Input_Signature_sample_quality_summary.tsv", "Input_Signature_classification_summary.tsv", "Input_Signature_analysis_definitions.tsv", "Input_Preop_platinum_signature_tests.tsv"])

    heading("Supplementary Figure 5. 요청된 환자별 signature 지표의 산출 가능성", 2)
    image("Supplementary_Figure_Signature_Requested_Metrics.png", 7.0)
    not_estimable = [r for r in signature_status if r["status"] == "Not estimable"]
    para(f"Platinum-associated, HRD-associated, MMR-associated process 각각에 대해 요청된 positive rate, 환자별 burden, 환자별 proportion의 총 9개 결과를 명시적으로 점검했다. 추정 불가 항목은 {len(not_estimable)}/9개였다. 각 항목의 No/Yes 품질평가 가능 환자 수가 모두 0이므로 통계검정과 효과크기를 계산하지 않고 NE로 표시했다.")
    para("NE는 signature가 생물학적으로 존재하지 않는다는 뜻이 아니라, 사전지정 품질기준 아래 현재 자료로 해당 환자별 지표를 신뢰성 있게 추정할 수 없다는 뜻이다. 관찰 후 cosine 또는 변이수 기준을 낮추면 양성률이 만들어질 수 있으나 post-hoc thresholding에 해당하므로 시행하지 않았다.")
    inputs(["Input_Signature_requested_metrics_status.tsv", "Input_Preop_platinum_signature_tests.tsv", "Input_Signature_analysis_definitions.tsv"])

    heading("Supplementary Figure 6. 집단 합산 hypothesis-directed SBS 민감도 분석", 2)
    image("Supplementary_Figure_Platinum_HRD_MMR_signatures.png", 6.1)
    plat_prop = lookup(signature_pooled, process="Platinum-associated SBS", metric="Pooled signature proportion")
    hrd_prop = lookup(signature_pooled, process="HRD-associated SBS", metric="Pooled signature proportion")
    mmr_prop = lookup(signature_pooled, process="MMR-associated SBS", metric="Pooled signature proportion")
    para(f"환자별 분석이 불가능하여 Yes와 No 각 군의 SBS96을 합산하고, standard pooled assignment에서 선택된 background signature에 SBS31/35, SBS3 및 MMR SBS를 추가한 constrained refit을 민감도 분석으로 시행하였다. 환자 단위 bootstrap 1,000회로 불확실성을, exposure label permutation 1,000회로 군 차이를 평가하고 6개 검정을 Holm 보정하였다.")
    para(f"Platinum-associated pooled proportion은 No {num(plat_prop['no_value']):.3f}, Yes {num(plat_prop['yes_value']):.3f}, permutation p={pval(plat_prop['permutation_p'])}, Holm p={pval(plat_prop['p_holm_6_tests'])}; HRD-associated는 No {num(hrd_prop['no_value']):.3f}, Yes {num(hrd_prop['yes_value']):.3f}, p={pval(hrd_prop['permutation_p'])}, Holm p={pval(hrd_prop['p_holm_6_tests'])}; MMR-associated는 두 군 모두 {num(mmr_prop['no_value']):.3f}, Holm p={pval(mmr_prop['p_holm_6_tests'])}였다. 어떤 process도 군 차이를 지지하지 않았다.")
    para("이 분석은 환자별 signature-positive 비율이 아니며 pooled spectrum의 민감도 분석이다. 또한 standard SigProfiler pooled assignment에서는 세 target process의 assigned count가 모두 0이었고, 그림의 값은 target signature를 포함시킨 constrained refit에서만 나온다. 따라서 논문의 주요 생물학적 결론이나 platinum 치료효과 근거로 사용하지 않고 방법론적 보충자료로만 제시한다. 모든 검체가 선행치료 전이므로 SBS31/35가 검출되더라도 현재 Preop_platinum_exposure에 의해 유발되었다고 해석할 수 없다.")
    inputs(["Input_Preop_platinum_pooled_signature_sensitivity.tsv", "Input_Signature_analysis_definitions.tsv", "Input_MAF_signature_audit.tsv"])

    heading("Supplementary Figure 7. 개별 SBS31/35, SBS3 및 MMR SBS 민감도 분석", 2)
    image("Supplementary_Figure_Individual_SBS_Pooled_Sensitivity.png", 7.1)
    individual_significant = [r for r in signature_individual if num(r.get("p_holm_20_tests")) < .05]
    proportion_rows = [r for r in signature_individual if r["metric"] == "Pooled signature proportion"]
    largest = max(proportion_rows, key=lambda r: abs(num(r["difference_yes_minus_no"])))
    para(f"통합 process뿐 아니라 platinum 관련 SBS31·SBS35, HRD 관련 SBS3, MMR 관련 SBS6·14·15·20·21·26·44를 각각 분리하였다. 각 signature에 대해 pooled attributed SNVs/patient와 pooled proportion을 비교하고 총 20개 permutation 검정을 Holm 보정했다. 보정 p<0.05인 결과는 {len(individual_significant)}개였다.")
    para(f"가장 큰 absolute proportion 차이는 {largest['signature']}에서 Yes-No={num(largest['difference_yes_minus_no']):.3f}였고, permutation p={pval(largest['permutation_p'])}, Holm p={pval(largest['p_holm_20_tests'])}였다. 개별 SBS 분해 역시 forced inclusion을 사용한 pooled sensitivity 결과이므로 환자 양성 판정이나 치료유발 기전의 증거로 사용하지 않는다.")
    inputs(["Input_Preop_platinum_pooled_signature_by_SBS.tsv", "Input_Signature_analysis_definitions.tsv"])

    heading("Supplementary Figure 8. Standard assignment와 constrained refit 감사", 2)
    image("Supplementary_Figure_Signature_Assignment_Audit.png", 6.8)
    standard_total = sum(num(r["no_standard_assignment_count"]) + num(r["yes_standard_assignment_count"]) for r in signature_assignment)
    constrained_total = sum(num(r["no_constrained_refit_count"]) + num(r["yes_constrained_refit_count"]) for r in signature_assignment)
    para(f"개별 target SBS의 standard pooled assignment와 target을 강제로 후보에 넣은 constrained refit의 attributed count를 나란히 제시했다. Target SBS 전체의 standard assigned count 합은 {standard_total:.1f}, constrained-refit 합은 {constrained_total:.1f}였다. 두 방법의 차이는 target 신호가 자료 주도 표준선택에서는 선택되지 않았고 가설지향 강제 적합에서만 분배되었음을 보여준다.")
    para("따라서 constrained 결과가 0보다 크다는 사실만으로 signature 존재를 확정할 수 없다. 본 분석의 적절한 결론은 '환자별 분석은 불가능하며, pooled sensitivity에서도 노출군 차이를 지지하는 견고한 근거가 없다'이다.")
    inputs(["Input_Signature_target_assignment_audit.tsv", "Input_Preop_platinum_pooled_signature_by_SBS.tsv"])

    heading("고찰: 임상적 의의, novelty와 강점", 1, page=True)
    heading("1. 규모보다 중요한 임상적 기여", 2)
    para("본 코호트의 1,011명은 수술 병리와 생존, 재발양상, targeted sequencing을 함께 보유한 단일 분석으로서 높은 통계적 정밀도와 아형별 비교 가능성을 제공한다. 그러나 novelty는 단순한 환자수 자체가 아니라, 동일한 환자에서 분자 아형–병리 표현형–치료 맥락–재발 후 경과를 연결했다는 점에 있다.")
    heading("2. 재발양상을 임상적으로 의미 있는 결과로 승격", 2)
    para("기존의 재발 유무 또는 DFS만으로는 재발 뒤의 질적 이질성을 설명하지 못한다. 본 결과는 Local + distant 재발이 기존 3군 분석에서 불량한 재발 후 생존과 연관되고, Distant only 내부에서는 Lung only가 Liver only 및 Other distant보다 양호함을 보여준다. 이는 재발환자 위험층화, 추적검사 이후 상담, 구제치료 연구의 층화변수 후보라는 임상적 의미가 있다. 다만 재발 후에 정의되므로 진단 시점의 예측모델 성능은 별도 검증이 필요하다.")
    heading("3. KRAS 이분형보다 아형 수준의 예후정보", 2)
    para("주요 driver의 단순 검출 여부는 다중보정 후 제한적인 독립 예후정보를 보였지만, KRAS 아형에서는 G12D와 G12V/G12R의 차이가 유지되었다. 이는 PDAC의 대표 driver를 단순 양성/음성으로 축약하지 않고 아형 수준으로 분석해야 한다는 임상·생물학적 근거를 강화한다.")
    heading("4. 선행 platinum 결과의 올바른 위치", 2)
    para("platinum 노출 분석은 임상적으로 중요한 질문이지만 이 관찰 코호트에서는 treatment indication, regimen, response 정보가 부족하고 구조적으로 선행치료 여부와 결합되어 있다. 유의하지 않은 결과를 '효과 없음'으로 단정하기보다, 치료 선택 편향과 불확실성을 정량화한 음성/탐색 결과로 제시하는 것이 타당하다.")
    heading("5. 제한점", 2)
    para("후향적 단일기관 자료, 검사범위·TMB 측정 세부사항 미확인, 치료정보의 제한, 결측, 다중검정, 수술환자 선택편향이 있다. 재발 분석은 기록된 재발을 사용하며 재발 전 사망을 포함하는 표준 DFS와 동일하지 않다. 원격재발 부위에는 전이부담과 구제치료의 잔여교란이 있다. CNV/LOH, tumor purity/cellularity, germline 및 기능적 HRD/dMMR는 분석하지 않았다. Targeted-panel MAF의 낮은 SNV 수와 BED 부재로 환자별 mutational signature를 판정할 수 없었다. 외부검증 전에는 결과를 치료의사결정 도구로 사용해서는 안 된다.")
    heading("최종 Summary", 2)
    para("대규모 PDAC 코호트에서 TP53–분화도와 CDKN2A–림프절 병기의 분자–병리 연관성을 확인했고, KRAS G12D의 상대적으로 불량한 예후 이질성을 보였다. 재발 분석에서는 기존 3군의 Local + distant 고위험뿐 아니라 Distant only 내부의 Lung-only 양호 경과가 추가로 확인되어 원격재발 부위의 임상적 이질성을 제시했다. 선행 platinum에 따른 HRD/MMR retained-variant 차이는 없었고, mutational signature는 자료 품질상 환자별 추론이 불가능했다. 따라서 논문의 중심 주장은 '분자 아형과 재발 부위별 표현형을 결합한 대규모 PDAC 임상경과 지도'이며, 치료효과와 signature 기전 주장은 배제하는 것이 논리적으로 가장 강하다.")

    heading("재현 방법과 입력", 1, page=True)
    para(f"저장소 루트에서 `Rscript --vanilla run_all.R --config=config/local.R`을 실행한다. local config의 clinical_file은 data/raw/260927_v3_PDAC_ANALYSIS_with_MAF_annotations.xlsx이고 KRAS 분석열은 KRAS_subtype_MAF이다. 분석 산출물은 outputs/updated_v3_20260928/{root.name}에 저장된다. Signature 분석은 local config에서 명시적으로 활성화하며, 별도로 설치된 SigProfiler 환경과 GRCh37 reference를 사용한다. 환자 단위 자료는 private/에만 있으며 배포하지 않는다.")
    inputs(["260927_v3_PDAC_ANALYSIS_with_MAF_annotations.xlsx", "PDAC_oncopanel.maf", "config/local.R", "scripts/manuscript_analysis.R", "R/manuscript_story_extensions.R", "R/recurrence_signature_extensions.R", "scripts/mutational_signature_analysis.py"])

    target.parent.mkdir(parents=True, exist_ok=True)
    doc.core_properties.title = "Updated PDAC clinicogenomic manuscript report"
    doc.core_properties.author = "PAAD research team"
    doc.save(target)


if __name__ == "__main__":
    ap = argparse.ArgumentParser()
    ap.add_argument("--results", required=True, type=Path)
    ap.add_argument("--output", required=True, type=Path)
    a = ap.parse_args()
    build_report(a.results.resolve(), a.output.resolve())
    print(a.output.resolve())
