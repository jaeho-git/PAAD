#!/usr/bin/env python3
"""Build the Korean manuscript report and shared presentation content.

Usage: python scripts/build_manuscript_documents.py --results outputs/.../manuscript
Reads aggregate publication tables only, plus run-level provenance. Patient-level
records are never embedded in Word, slides, or content.json.
"""
import argparse
import csv
import json
import math
from pathlib import Path
from datetime import date
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from PIL import Image


def build_content(root):
    def read(name):
        with (root / "tables" / name).open(encoding="utf-8", newline="") as handle:
            return list(csv.DictReader(handle, delimiter="\t"))
    def num(x):
        try: return float(x)
        except (TypeError, ValueError): return float("nan")
    def f(x, digits=2):
        n = num(x)
        return f"{n:.{digits}f}" if math.isfinite(n) else "NE"
    def p(x):
        n = num(x)
        return "NE" if not math.isfinite(n) else "<0.001" if n < .001 else f"{n:.3f}"
    def hr(r): return f"{f(r['HR'])} (95% CI {f(r['lower95'])}–{f(r['upper95'])})"
    prov = json.loads((root / "private" / "provenance.json").read_text())
    flow = {r["item"]: int(r["n"]) for r in read("Input_Cohort_flow.tsv")}
    cox = read("Supplementary_Table3_Cox_coefficients.tsv")
    core = [r for r in cox if r["family"] == "Extended genomic" and r["is_exposure"] == "TRUE"]
    ph = read("Supplementary_Table3_PH_tests.tsv")
    pairs = [r for r in read("Supplementary_Table4_Cox_pairs.tsv") if r["exposure"] == "KRAS_group" and r["family"] == "Extended genomic" and r["endpoint"] == "OS"]
    med = read("Supplementary_Table4_KM_estimates.tsv")
    km = [r for r in med if r["variable"] == "KRAS_group" and r["endpoint"] == "OS"]
    cmh = read("Supplementary_Table2_Pathology_CMH.tsv")
    lookup = lambda rows, **kw: next(r for r in rows if all(r[k] == value for k, value in kw.items()))
    kras = lookup(core, exposure="KRAS")
    tp = lookup(cmh, gene="TP53", characteristic="Differentiation")
    cd = lookup(cmh, gene="CDKN2A", characteristic="N_stage")
    dr = lookup(pairs, group1="G12D", group2="G12R")
    dv = lookup(pairs, group1="G12D", group2="G12V")
    tmb = read("Supplementary_Table7_Reported_TMB.tsv")[0]
    tc = lookup(cox, exposure="TMB_SD", is_exposure="TRUE")
    tphp = lookup(ph, exposure="TMB_SD", term="TMB_SD")["p"]
    krph = lookup(ph, exposure="KRAS_group", family="Extended genomic", term="KRAS_group")["p"]
    rrph = lookup(ph, exposure="KRAS_group", family="Adjusted recurrence", term="KRAS_group")["p"]
    source = "임상정보 KRAS_subtype" if prov["kras_source"] == "clinical" else "MAF에서 도출한 KRAS subtype"
    refs = [
        {"label": "McIntyre CA et al. Cancer. 2020;126:3939–3949.", "url": "https://doi.org/10.1002/cncr.33038"},
        {"label": "Campbell BA et al. Ann Surg. 2025;282:525–533.", "url": "https://doi.org/10.1097/SLA.0000000000006794"},
        {"label": "Varghese AM et al. Nat Med. 2025;31:466–477.", "url": "https://doi.org/10.1038/s41591-024-03362-3"},
        {"label": "Ardalan B et al. Clin Cancer Res. 2025;31:1082–1090.", "url": "https://doi.org/10.1158/1078-0432.CCR-24-3149"},
        {"label": "Merino DM et al. J Immunother Cancer. 2020;8:e000147.", "url": "https://doi.org/10.1136/jitc-2019-000147"},
        {"label": "Altman DG et al. REMARK explanation and elaboration. PLoS Med. 2012;9:e1001216.", "url": "https://doi.org/10.1371/journal.pmed.1001216"},
    ]
    methods = [
        f"업데이트된 임상자료 {flow['Source records']:,}행 중 중복 patient_id에 해당하는 {flow['Duplicate-ID records excluded']}행을 모두 제외하였다. 분석 단위는 환자이며 총 {flow['Analysis patients']:,}명이다. M0 {flow['M0']}명과 M1 {flow['M1']}명을 포함하였다. 결과별 결측은 해당 분석에서만 제외하였다.",
        f"KRAS 아형은 {source} 하나를 사용하였다. 병리·생존분석의 KRAS 변이 유무와 주요 변이 유전자 수도 이 정의를 유지하였다. Oncoplot의 유전자 행은 KRAS를 포함하여 모두 MAF의 실제 변이와 유형을 표시한다. 상단 KRAS 아형 annotation은 선택한 아형 출처를 사용한다. 검출되지 않은 상태가 확인된 생물학적 wild type을 의미하지는 않는다.",
        "Oncoplot의 상단은 환자별 전체 유지 변이 수, 오른쪽은 해당 유전자의 변이 수, 왼쪽은 변이 보유 환자 수와 표시 코호트 내 비율이다. 환자·유전체 위치·대립유전자로 동일 변이를 식별하여 중복 annotation을 제거하였다. 같은 변이가 다른 환자에게 있으면 각각 집계하였다. 유전자별 검사 범위는 확인되지 않아 비율을 전체 환자의 생물학적 유병률로 해석하지 않는다.",
        "OS는 임상자료에 기록된 수술 기준 추적기간과 사망 여부를 사용하였다. KM 곡선에는 검열 표시와 number at risk를 제시하였다. 전체 log-rank와 쌍별 log-rank를 구분하고, 모든 쌍별 비교에 Holm 보정을 적용하였다. 기록된 재발은 사망을 사건에 포함하는 DFS/RFS와 구분하였다.",
        "Cox 기본 모형은 연령(자연 스플라인 3 자유도), 성별, T category, 선행치료, M category를 보정하였다. 본문의 병리 확장 모형은 N category, 분화도, LVI, PNI, 절제연을 추가하였다. 동일 보정 수준의 유전체 노출 계수 전체에 BH 보정을 적용하였다. 계수의 95% CI는 개별 구간이며 동시 신뢰구간은 아니다. 모형별 분석 인원, 사건 수, 추정 상태와 비례위험 검정을 함께 보고하였다.",
        "유전자–병리 연관성은 선행치료 여부로 층화한 generalized CMH로 평가하였다. 4개 유전자와 6개 병리 항목의 총 24개 검정에 BH 보정을 적용하였다. 다범주 병리 항목의 쌍별 검정은 항목별 Holm 보정을 적용하였다. Table 1과 Table 2는 기술통계표이며 유의성 검정에 따른 변수 선택에 사용하지 않았다.",
        "Reported TMB는 임상자료의 원래 값을 사용하였다. 단위·계산법·검사 버전 및 결측 사유를 확인할 수 없어 연속형 탐색 분석으로 제한하였다. 외부 TMB-high 기준을 적용하지 않았고 결측값을 대체하지 않았다. Variant count는 oncoplot의 기술적 정보로 사용하였다.",
    ]
    figs = []
    def fig(identifier, title, image, assets, method, finding, interpretation, limitation, inputs):
        figs.append(dict(id=identifier,title=title,image=image,assets=assets,method=method,finding=finding,interpretation=interpretation,limitation=limitation,inputs=inputs))
    fig("Figure 1", "연구 코호트와 PDAC의 유전체 특성", "Figure1_Cohort_and_genomics.png",
        ["Figure1A_Cohort.png","Figure1B_Oncoplot.png","Figure1C_KRAS_distribution.png"],
        f"Figure 1a는 환자 선정 흐름도, 1b는 변이·임상정보 oncoplot, 1c는 KRAS 아형 분포이다. Oncoplot의 각 열은 환자이며 KRAS를 포함한 모든 유전자 행은 MAF의 변이 유형을 표시한다. 상단 KRAS 아형은 {source}를 사용한다. 상단 막대는 환자별 전체 Variant count, 오른쪽은 유전자별 변이 수, 왼쪽은 변이 보유 환자 n (%)이다. 임상정보 16개를 동일한 환자 순서로 표시하며 전체 항목과 결측 수는 임상 track 입력표에 제시하였다.",
        f"분석 대상은 {flow['Analysis patients']:,}명이며 사망 {flow['OS deaths']}건을 관찰하였다. KRAS 아형별 환자 수는 " + ", ".join(f"{r['group']} {int(float(r['n']))}명" for r in km) + "이다.",
        "이 그림은 이후 병리 및 생존분석의 모집단과 유전체 구성을 정의한다. 유전자별 변이 개수와 변이 보유 환자 수를 함께 제시하여 한 환자의 다중 변이가 빈도 해석에 미치는 영향을 구분할 수 있다. M1도 분석에 포함했지만 M0가 대부분인 코호트임을 전제로 해석한다.",
        "빈도가 높다는 사실만으로 예후 영향이나 치료 표적성을 판단할 수 없다. 변이의 병원성, 생식세포 기원 및 기능적 결핍을 확정하지 않았다. 검사 범위가 확인되지 않은 유전자의 회색 칸은 유지된 변이가 없다는 의미로만 읽는다.",
        ["Input_Cohort_flow.tsv","Input_KRAS_distribution.tsv","Input_Figure1B_Oncoplot_gene_counts.tsv","Input_Figure1B_Oncoplot_clinical_tracks.tsv"])
    figs[-1]['panels'] = [
        dict(id="Figure 1a",title="분석 대상 환자의 선정",image="Figure1A_Cohort.png",
             caption=f"전체 {flow['Source records']:,}개 기록에서 중복 ID를 가진 {flow['Duplicate patient IDs excluded']}명의 기록 {flow['Duplicate-ID records excluded']}개를 모두 제외하였다. 최종 {flow['Analysis patients']:,}명에서 M0 {flow['M0']}명과 M1 {flow['M1']}명을 모두 포함한다. 상자는 각 단계의 모집단, 화살표는 제외 및 포함 관계를 나타낸다. M0와 M1은 최종 코호트의 하위집단이며 추가 제외 기준이 아니다.",inputs=["Input_Cohort_flow.tsv"],diagram=json.loads((root/'private/cohort_flow_diagram.json').read_text())),
        dict(id="Figure 1b",title="환자별 유전자 변이와 임상정보",image="Figure1B_Oncoplot.png",
             caption=f"KRAS를 포함한 모든 유전자 행은 MAF의 실제 변이 유형을 표시하고 상단 KRAS 아형은 {source}를 사용한다. 환자별 전체 Variant count와 유전자별 변이 수를 임상정보 16개와 함께 표시하였다. 임상정보는 KRAS 아형, 성별, 연령군, 분화도, 선행치료, T·N·M category, 림프절 양성 여부, BMI, CA19-9, CEA, AJCC 병기, 병기군, 종양 크기, Reported TMB이다. 임상정보의 회색은 결측이다. CA19-9와 CEA의 색은 80백분위수에서 포화되며 값 자체는 변경하지 않았다. TMB와 종양 크기의 단위는 원자료 기준으로 미확인이다.",inputs=["Input_Figure1B_Oncoplot_gene_counts.tsv","Input_Figure1B_Oncoplot_clinical_tracks.tsv"]),
        dict(id="Figure 1c",title="KRAS 아형별 환자 분포",image="Figure1C_KRAS_distribution.png",
             caption="선택한 단일 KRAS 출처를 기준으로 각 아형의 환자 수와 전체 분석 환자 중 비율을 표시하였다.",inputs=["Input_KRAS_distribution.tsv"]),
    ]
    fig("Figure 2", "주요 유전자와 임상병리학적 특성의 연관성", "Figure2_Genes_and_pathology.png",
        ["Figure2A_Pathology_matrix.png","Figure2B_TP53_grade.png","Figure2C_CDKN2A_nodes.png"],
        "전체 연관성 행렬에는 24개 generalized CMH 검정의 p와 BH q를 표시하였다. 선행치료 여부로 층화했으며, 상세 패널은 TP53과 분화도, CDKN2A와 N category의 분포를 보여준다. 막대의 백분율은 해당 유전자 상태 집단 내 비율이다.",
        f"TP53–분화도는 p {p(tp['p'])}, q {p(tp['q_BH'])}이고 CDKN2A–N category는 p {p(cd['p'])}, q {p(cd['q_BH'])}이다. 이는 여러 항목을 동시에 검토한 뒤에도 남는 연관성이다. 상세한 범주별 환자 수와 전체 쌍별 결과는 Supplementary Table 2의 입력표에 보존하였다.",
        "주요 유전자 변이가 동일한 병리 소견과 일괄적으로 연결되는 것은 아니다. 분화도와 림프절 병기의 관찰된 연관성은 유전체와 병리 표현형을 함께 설명하는 근거이다. 상세 분포와 효과 방향을 함께 읽어야 하며, 작은 p값만으로 임상적 차이의 크기를 판단하지 않는다.",
        "선행치료 층화는 모든 교란요인을 보정한 분석이 아니다. 수술 후 병리와 치료 후 병리는 치료 이전의 종양 상태와 다를 수 있다. 이 결과로 변이가 저분화 또는 전이를 직접 유발한다고 결론 내릴 수 없다.",
        ["Supplementary_Table2_Pathology_CMH.tsv","Input_Pathology_counts.tsv","Supplementary_Table2_Pathology_pairs.tsv"])
    gene_find = "각 유전자의 비교군은 변이 검출군, 기준군은 미검출군이다. " + " ".join(f"{g}: HR {hr(lookup(core,exposure=g))}, q {p(lookup(core,exposure=g)['q_BH'])}." for g in ["KRAS","TP53","SMAD4","CDKN2A"])
    count_rows = [r for r in core if r["exposure"] == "driver_group"]
    count_sig = any(num(r["q_BH"]) < .05 for r in count_rows)
    fig("Figure 3", "주요 유전자 및 변이 유전자 수와 전체생존", "Figure3_Gene_level_OS.png",
        ["Figure3A_Gene_OS.png","Figure3B_Driver_count_OS.png","Figure3C_Driver_count_adjusted.png"],
        "각 유전자를 별도로 투입한 Cox 모형에서 연령 스플라인, 성별, T/N/M category, 선행치료, 분화도, LVI/PNI, 절제연을 보정하였다. 이는 다른 유전자 모두를 동시에 보정한 모형과 구분한다. 주요 변이 유전자 수는 KRAS·TP53·SMAD4·CDKN2A 중 변이가 있는 유전자의 개수이다. 0, 1, 2, 3개 이상 집단을 비교하며 기준군은 1개이다. HR>1이면 표시된 비교군의 사망 hazard가 기준군보다 높다.",
        gene_find + (" 변이 유전자 수 비교 중 일부가 BH 보정 후 유의하다." if count_sig else " 변이 유전자 수의 각 비교는 BH 보정 후 유의하지 않았다."),
        "병리 특성과의 연관성과 전체생존과의 연관성은 별개이다. 보정 후 효과크기와 신뢰구간을 확인하면 각 유전자에 대한 예후 근거의 범위를 구분할 수 있다. 변이 유전자 수에 따른 단순한 용량–반응 관계는 별도 검정 없이 주장하지 않는다.",
        "관찰연구의 HR은 치료 효과나 인과효과를 나타내지 않는다. 한 유전자의 p값이 유의하고 다른 유전자는 유의하지 않다는 사실만으로 두 효과가 통계적으로 다르다고 볼 수 없다. 작은 집단과 후기 위험집단 감소, 측정되지 않은 치료·선택 요인이 불확실성을 남긴다.",
        ["Supplementary_Table3_Cox_coefficients.tsv","Supplementary_Table4_Logrank.tsv","Supplementary_Table4_Logrank_pairs.tsv","Supplementary_Table3_PH_tests.tsv"])
    medtext = ", ".join(f"{r['group']} {f(r['median'],1)}개월" for r in km if r["group"] in ["G12D","G12V","G12R"])
    fig("Figure 4", "KRAS 아형에 따른 전체생존", "Figure4_KRAS_subtype_OS.png",
        ["Figure4A_KRAS_OS.png","Figure4B_KRAS_pairwise.png"],
        "KRAS 5개 범주의 전체생존을 KM과 log-rank로 평가하고, 동일 임상 변수로 보정한 Cox 모형에서 10개 아형 쌍을 비교하였다. 곡선의 유의한 쌍은 비보정 생존의 Holm 보정 log-rank이며, forest plot은 보정 Cox 대비이다. 두 p값은 서로 다른 질문에 답한다.",
        f"중앙 OS는 {medtext}이다. G12R을 기준으로 G12D의 보정 사망 HR은 {hr(dr)}, Holm p {p(dr['p_holm'])}이다. G12V를 기준으로 G12D의 HR은 {hr(dv)}, Holm p {p(dv['p_holm'])}이다. KRAS 아형의 OS 비례위험 검정 p는 {p(krph)}이다.",
        "KRAS 변이 환자를 하나의 집단으로 묶으면 아형별 예후 차이를 충분히 표현하지 못할 수 있다. 아형 구분은 향후 예후 연구와 임상시험에서 층화 변수로 검토할 근거가 된다. 다만 여기서 보이는 중앙생존 차이를 아형에 따른 치료 이득의 크기로 바꾸어 해석하지 않는다.",
        "다중비교 보정 후 결과와 HR의 방향을 함께 읽는다. Not detected는 기록 또는 분석 기준에서 검출되지 않은 상태이며 확정된 생물학적 wild type과 동일하지 않다. Other rare/multiple은 희귀·다중 아형의 묶음으로, G12D를 제외한 모든 변이 환자를 뜻하는 Non-G12D mutant와 구분한다.",
        ["Supplementary_Table4_KM_estimates.tsv","Supplementary_Table4_Cox_pairs.tsv","Supplementary_Table4_Logrank_pairs.tsv","Supplementary_Table4_Number_at_risk.tsv"])
    for s, number in [("HRD",1),("MMR",2)]:
        counts = read(f"Input_Supplementary_Figure{number}_{s}_gene_counts.tsv")
        tops = sorted(counts,key=lambda r:int(r["patients"]),reverse=True)[:3]
        finding = "주요 변이 보유 환자 수는 " + ", ".join(f"{r['gene']} {r['patients']}명 ({f(r['percent'],1)}%)" for r in tops) + "이다. 각 유전자 행의 변이 개수와 환자 수를 따로 표시했으며, 동일 환자가 여러 유전자에 기여할 수 있다."
        fig(f"Supplementary Figure {number}",f"{'HRD 관련' if s=='HRD' else 'MMR'} 유전자 변이와 임상정보",f"Supplementary_Figure{number}_{s}.png",[f"Supplementary_Figure{number}_{s}.png"],
            f"McIntyre 연구의 {'18개 HRD 관련' if s=='HRD' else '4개 MMR'} 유전자 목록을 사용하였다. 전체 분석 환자를 유지하고 임상정보 12개를 같은 환자 순서로 표시하였다. KRAS 아형, 연령, 성별, 선행치료, Preop_platinum, T·N·M category, 분화도, 절제연, MSI, Reported TMB가 포함된다. 환자별 총 Variant count는 그림에 포함된 유전자군에만 한정하지 않는다.",finding,
            "이 그림은 DNA 복구 관련 유전자에서 어떤 변이가 어떤 환자에 분포하는지 탐색하는 자료이다. 임상 annotation은 변이 보유 환자의 병기·선행치료·분화도 구성을 함께 살펴보게 한다. 시각적으로 비슷하게 모여 보이는 현상만으로 통계적 연관성을 확정하지 않는다.",
            "유전자 목록의 변이 보유는 기능적 HRD 또는 dMMR 진단이 아니다. 모든 변이의 병원성을 확인한 것은 아니며, 치료 반응·PARP 억제제 감수성·면역치료 효과를 추정하지 않는다. 변이 개수가 많은 환자는 목록 내 변이가 발견될 기회도 늘 수 있다.",
            [f"Input_Supplementary_Figure{number}_{s}_gene_counts.tsv",f"Input_Supplementary_Figure{number}_{s}_clinical_tracks.tsv","Supplementary_Table5_Repair_gene_lists.tsv","Supplementary_Table5_Repair_clinical_profiles.tsv"])
    rec = read("Supplementary_Table6_Recurrence_patterns.tsv")
    fig("Supplementary Figure 3","기록된 재발과 재발 양상","Supplementary_Figure3_Recurrence.png",["Supplementary_Figure3_Recurrence.png"],
        "기록된 재발 여부와 추적기간이 있는 환자에서 KRAS 아형별 재발까지의 시간을 탐색하였다. 재발 양상 막대는 재발이 기록된 환자들 내 분포이며 전체 코호트의 장기별 누적발생률이 아니다.",
        f"재발 분석은 {flow['Recurrence evaluable']}명, 사건 {flow['Recorded recurrences']}건이다. " + ", ".join(f"{r['Recurrence_pattern']} {r['n']}명 ({f(r['percent'],1)}%)" for r in rec) + f"이다. KRAS 아형의 비례위험 검정 p는 {p(rrph)}이다.",
        "재발 결과는 OS 결과를 보완하는 탐색 자료이다. 재발이 발생한 환자 중 원격·국소 양상의 분포를 기술할 수 있지만, 각 아형의 특정 장기 전이 성향이나 전체 환자의 재발 위험을 이 비율로 직접 추론하지 않는다.",
        "M1에서는 이 재발 추적변수가 정의되지 않아 해당 endpoint에서만 빠진다. 사망 포함 DFS/RFS 및 사망을 경쟁사건으로 처리한 누적발생률과 구분한다. 비례위험 가정 위반이 있는 모형의 단일 HR은 시간 전체의 일정한 효과로 해석하지 않는다.",
        ["Supplementary_Table6_Recurrence_patterns.tsv","Supplementary_Table4_KM_estimates.tsv","Supplementary_Table3_PH_tests.tsv"])
    fig("Supplementary Figure 4","Reported TMB의 탐색적 분석","Supplementary_Figure4_Reported_TMB.png",["Supplementary_Figure4_Reported_TMB.png"],
        "임상자료에 기록된 TMB의 전체 분포와 KRAS 아형별 분포를 제시하였다. 다범주 비교는 Kruskal–Wallis 및 Dunn 쌍별 검정을 사용하였다. 관련 6개 전체 검정에는 BH, 각 변수의 쌍별 비교에는 Holm 보정을 적용하였다. OS 모형에서는 관측 TMB의 1 SD를 단위로 사용하였다.",
        f"TMB는 {tmb['observed']}명에서 기록되었고 {tmb['missing']}명에서 결측이다. 중앙값은 {f(tmb['median'],1)} [IQR {f(tmb['q1'],1)}–{f(tmb['q3'],1)}]이다. OS의 1 SD당 보정 HR은 {hr(tc)}, p {p(tc['p'])}이며 비례위험 검정 p는 {p(tphp)}이다.",
        "Reported TMB의 분포를 임상적 맥락에서 기술할 수 있으나, 검사법과 단위가 확인되지 않은 상태의 탐색 결과이다. 유의하지 않은 생존 결과만으로 예후와의 연관성이 없다고 확정할 수 없다. 본 논문의 핵심 예후 주장은 KRAS 및 주요 유전자 분석에 둔다.",
        "단위·계산법·검사 버전·결측 원인을 확인하지 못했다. 외부 TMB-high 기준이나 면역치료 적응 기준을 적용하지 않았다. 결측은 0이나 Variant count로 대체하지 않았으며, 관측 환자만의 분석에 선택 편향이 남을 수 있다. PH 가정 위반이 있으면 일정한 HR 해석을 제한한다.",
        ["Supplementary_Table7_Reported_TMB.tsv","Supplementary_Table7_TMB_associations.tsv","Supplementary_Table7_TMB_pairs.tsv","Supplementary_Table3_Cox_coefficients.tsv"])
    fig("Supplementary Figure 5","보정 모형에 따른 주요 유전자 효과 추정","Supplementary_Figure5_Model_sensitivity.png",["Supplementary_Figure5_Model_sensitivity.png"],
        "핵심 보정 모형과 N category, LVI, PNI, 절제연을 추가한 확장 모형에서 주요 유전자 HR을 비교하였다. 각 모형은 이용 가능한 공변량을 가진 환자를 대상으로 하며, 분석 인원과 사건 수는 진단표에 제시하였다.",
        "각 유전자에 대해 두 모형의 HR과 95% CI를 나란히 제시하였다. 효과 방향·크기와 불확실성의 변화를 읽어야 하며, 유의성 여부만으로 모형 간 차이를 판정하지 않는다.",
        "측정한 병리 변수를 추가했을 때 추정치가 얼마나 달라지는지를 보여주는 보조 분석이다. 큰 변화가 있다면 교란 또는 조정 변수의 역할을 검토할 수 있다. 변화가 작더라도 측정되지 않은 교란이 없다는 뜻은 아니다.",
        "병리 변수는 유전체와 예후 사이의 경로에 놓일 수도 있어, 확장 모형이 항상 더 정확한 인과효과를 주는 것은 아니다. 두 모형은 독립적인 검증 코호트가 아니며 외부 검증을 대체하지 않는다.",
        ["Supplementary_Table3_Cox_coefficients.tsv","Supplementary_Table3_Model_diagnostics.tsv","Supplementary_Table3_PH_tests.tsv"])

    labels = {"Age":"Age, years","Sex":"Sex","M_stage":"M category","Neoadjuvant":"Neoadjuvant treatment","Preop_platinum_exposure":"Preop_platinum","T_stage":"T category","N_stage":"N category","Differentiation":"Differentiation","LVI":"LVI","PNI":"PNI","R_status":"Margin","Adjuvant_treatment":"Adjuvant_Tx","TMB":"Reported TMB"}
    labels.update({'OS_m':'OS time, months','RFS_m':'Recorded recurrence time, months','Recur':'Recorded recurrence','KRAS_group':'KRAS subtype','Stage_Group':'Stage group','HRD_variant_status':'HRD-list variant status','MMR_variant_status':'MMR-list variant status'})
    def characteristic(r):
        cat = {"y":"Yes","n":"No","Median [IQR]":"median [IQR]"}.get(r["category"],r["category"])
        return labels.get(r["variable"],r["variable"]) + ": " + cat
    tables=[]
    def contrast_label(r):
        term=r['term'];ex=r['exposure']
        if ex in ['KRAS','TP53','SMAD4','CDKN2A']:return ex+' detected vs not detected'
        if ex=='KRAS_group':return term.removeprefix('KRAS_group')+' vs not detected'
        if ex=='driver_group':return 'Driver genes: '+term.removeprefix('driver_group')+' vs 1'
        return {'Any_driver':'Any driver vs none','TP53_truncationTruncating':'TP53 truncating vs nontruncating','KRAS_TP53One gene':'KRAS/TP53: one vs neither','KRAS_TP53Both genes':'KRAS/TP53: both vs neither','G12D_vs_otherG12D':'G12D vs non-G12D mutant'}.get(term,term)
    def table(identifier,title,headers,rows,inputs,note): tables.append(dict(id=identifier,title=title,headers=headers,rows=rows,inputs=inputs,note=note))
    base=read("Table1_Cohort_characteristics.tsv")
    table("Table 1","전체 코호트의 임상병리학적 특성",["Characteristic","Value","Missing n"],[[characteristic(r),r["value"],r["missing"]] for r in base],["Table1_Cohort_characteristics.tsv","Input_Followup.tsv"],"연속형은 중앙값 [IQR], 범주형은 전체 코호트 내 n (%)이다. 같은 변수의 결측 수는 각 범주 행에서 반복 표시하였다. Reported TMB의 단위는 미확인이다.")
    kkrows=read("Table2_KRAS_characteristics.tsv");groups=["Not detected","G12D","G12V","G12R","Other rare/multiple"]
    ordered=list(dict.fromkeys((r["variable"],r["category"]) for r in kkrows))
    rows=[]
    for nm,cat in ordered:
        rr=next(r for r in kkrows if r["variable"]==nm and r["category"]==cat)
        rows.append([characteristic(rr)]+[next(r["value"] for r in kkrows if r["variable"]==nm and r["category"]==cat and r["group"]==g) for g in groups])
    table("Table 2","KRAS 아형별 임상병리학적 특성",["Characteristic"]+[f"{g}\nn={int(float(lookup(km,group=g)['n']))}" for g in groups],rows,["Table2_KRAS_characteristics.tsv"],"범주형 비율의 분모는 각 KRAS 집단 전체이다. 결측으로 범주 합계가 100%보다 작을 수 있다. 아형 간 구성 차이를 설명하는 기술통계이며 치료 효과 비교표가 아니다.")
    table("Supplementary Table 1","분석 변수의 결측 현황",["Variable","Missing n","Total n"],[[labels.get(r["variable"],r["variable"]),r["missing_n"],r["total_n"]] for r in read("Supplementary_Table1_Missingness.tsv")],["Supplementary_Table1_Missingness.tsv"],"결측은 분석별로 처리하였다. 기록된 재발 변수의 결측에는 M1에서의 endpoint 미정의가 포함된다.")
    table("Supplementary Table 2","유전자와 병리 특성의 전체 검정",["Gene","Characteristic","N","p","BH q"],[[r["gene"],labels.get(r["characteristic"],r["characteristic"]),r["n"],p(r["p"]),p(r["q_BH"])] for r in cmh],["Supplementary_Table2_Pathology_CMH.tsv","Supplementary_Table2_Pathology_pairs.tsv"],"선행치료 층화 generalized CMH. BH 보정 범위는 24개 검정이다. 전체 쌍별 결과 및 방향을 정의한 세부 항목은 Input 2에 포함하였다.")
    table("Supplementary Table 3","주요 유전체 모형의 보정 효과 추정",["Exposure / contrast","N","HR [95% CI]","p","BH q"],[[contrast_label(r),r["n"],f"{f(r['HR'])} [{f(r['lower95'])}, {f(r['upper95'])}]",p(r["p"]),p(r["q_BH"])] for r in core],["Supplementary_Table3_Cox_coefficients.tsv","Supplementary_Table3_Model_diagnostics.tsv","Supplementary_Table3_PH_tests.tsv"],"핵심 보정 모형의 노출 계수 전체를 제시하였다. 공변량 계수와 확장 모형은 Input 1, 적합 상태는 Input 2, PH 검정은 Input 3에 포함하였다. q는 동일 보정 수준의 전체 유전체 노출 계수를 대상으로 계산하였다.")
    table("Supplementary Table 4","KRAS 아형의 보정 쌍별 생존 비교",["Group 1 vs group 2","HR [95% CI]","Raw p","Holm p"],[[r['group1']+" vs "+r['group2'],f"{f(r['HR'])} [{f(r['lower95'])}, {f(r['upper95'])}]",p(r['p']),p(r['p_holm'])] for r in pairs],["Supplementary_Table4_Cox_pairs.tsv","Supplementary_Table4_Logrank_pairs.tsv","Supplementary_Table4_KM_estimates.tsv"],"HR은 group 1 대 group 2이다. 10개 Cox 대비에 Holm 보정을 적용하였다. CI는 개별 95% 구간이며, 보정 p와 CI의 유의성 판단이 다를 수 있다. 비보정 log-rank는 별도 입력표에 보존하였다.")
    repairrows=[]
    for s,n in [("HRD",1),("MMR",2)]:
        for r in read(f"Input_Supplementary_Figure{n}_{s}_gene_counts.tsv"):
            repairrows.append([s,r['gene'],r['variant_count'],f"{r['patients']} ({f(r['percent'],1)}%)"])
    table("Supplementary Table 5","DNA 복구 관련 유전자의 변이 및 환자 빈도",["Gene list","Gene","Variants n","Patients n (%)"],repairrows,["Input_Supplementary_Figure1_HRD_gene_counts.tsv","Input_Supplementary_Figure2_MMR_gene_counts.tsv","Supplementary_Table5_Repair_clinical_profiles.tsv"],"비율의 분모는 표시 코호트 전체이다. 유전자별 검사 범위와 변이의 기능적 의미를 확인하지 못했으므로 HRD/dMMR 유병률로 해석하지 않는다. 임상 분포의 상세 값은 Input 3에 포함하였다.")
    table("Supplementary Table 6","기록된 재발의 양상",["Pattern","Patients n","Among recurrences (%)"],[[r['Recurrence_pattern'],r['n'],f(r['percent'],1)] for r in rec],["Supplementary_Table6_Recurrence_patterns.tsv","Supplementary_Table4_KM_estimates.tsv"],"전체 코호트의 누적 재발률이 아니라, 재발이 기록된 환자들 안에서의 구성비이다.")
    tr=read("Supplementary_Table7_TMB_associations.tsv")
    table("Supplementary Table 7","Reported TMB의 탐색적 연관성",["Grouping","Observed n","p","BH q"],[[labels.get(r['variable'],r['variable']),r['n'],p(r['p']),p(r['q_BH'])] for r in tr],["Supplementary_Table7_TMB_associations.tsv","Supplementary_Table7_TMB_pairs.tsv","Supplementary_Table7_Reported_TMB.tsv"],"TMB의 단위와 검사법은 미확인이다. 전체 검정은 Kruskal–Wallis이며 6개 검정에 BH 보정을 적용하였다. Dunn/Wilcoxon 쌍별 결과는 Input 2에 제시하였다.")
    discussion=[
        ("임상병리와 예후의 서로 다른 연관성",f"TP53–분화도와 CDKN2A–림프절 병기의 연관성은 병리 표현형과 유전체를 함께 설명한다. 생존에서는 KRAS 변이 유무 및 아형을 중심으로 효과크기를 평가하였다. 병리 연관성이 곧 독립적인 생존 영향이라는 전제를 두지 않은 점이 결과 해석의 핵심이다."),
        ("기존 PDAC 연구와의 관계","본 분석의 1,011명은 McIntyre 연구의 수술 코호트 283명과 Campbell 연구 508명보다 많다 [1,2]. 하지만 Varghese 연구는 전체 2,336명을 포함했으므로 최대 규모라는 주장은 적절하지 않다 [3]. 각 연구의 병기 구성, 검사 범위와 생존 시작점이 달라 절대 생존기간이나 변이 빈도를 단순 비교하지 않는다."),
        ("임상적 의의와 차별성","KRAS 아형별 예후 차이는 기존 연구에도 보고되어 있다 [4]. 본 연구의 기여는 상세 병리정보와 생존정보를 연결한 1,000명 이상 코호트에서 주요 유전자와 KRAS 아형의 연관성을 평가하는 데 있다. 환자 수 자체보다 분석 가능한 사건 수, 효과 추정의 정밀도, 임상적 맥락을 함께 제시하는 것이 강점이다. 아형은 후속 예후 연구나 임상시험 층화 변수의 후보이며, 이 자료만으로 치료를 선택하거나 변경할 근거는 아니다."),
        ("한계와 결과의 적용 범위","본 결과는 관찰된 연관성이다. 병기와 치료 배경의 불균형, NGS 검사 대상 선택, 수술·검체 채취·검사 시점의 관계 및 잔여 교란을 완전히 해결하지 못했다. Reported TMB의 측정 세부사항과 결측 원인, 유전자별 검사 가능 범위가 불명확하다. 희귀 KRAS 아형과 M1 하위집단의 추정은 제한적이며, 다른 기관의 환자에서 검증이 필요하다."),
        ("Summary",f"총 {flow['Analysis patients']:,}명의 PDAC 코호트에서 주요 유전체 특성을 임상병리 및 전체생존과 연결하였다. TP53–분화도와 CDKN2A–N category의 연관성을 관찰했고, KRAS 아형별 OS 차이를 임상 변수 보정 및 다중비교와 함께 평가하였다. 연구의 임상적 의미는 분자적 이질성에 따른 예후를 이해하는 데 있다. 기능적 HRD/dMMR, 치료 반응 예측 또는 즉각적인 진료 변경까지 입증한 결과로 확대하지 않는다.")
    ]
    content = dict(title="PDAC 임상유전체 특성과 생존 연관성",subtitle=f"{flow['Analysis patients']:,}명 코호트의 주요 유전자 및 KRAS 아형 분석",kras_source=prov['kras_source'],methods=methods,figures=figs,tables=tables,discussion=[dict(title=a,text=b) for a,b in discussion],references=refs)
    if prov.get('clinical_extensions'):
        from manuscript_clinical_interpretation import extend_content
        content = extend_content(content, root)
    return content


def build_docx(content, root, target):
    doc=Document();sec=doc.sections[0]
    sec.page_width=Inches(8.5);sec.page_height=Inches(11)
    sec.top_margin=sec.bottom_margin=Inches(.65);sec.left_margin=sec.right_margin=Inches(.7)
    for style in doc.styles:
        if style.type in (1,2):
            style.font.name="Pretendard";style.font.color.rgb=RGBColor(0,0,0)
            fonts=style.element.get_or_add_rPr().get_or_add_rFonts()
            for a in list(fonts.attrib):
                if "theme" in a.lower():del fonts.attrib[a]
            for key in ['ascii','hAnsi','eastAsia','cs']:fonts.set(qn('w:'+key),'Pretendard')
            for border in list(style.element.iter(qn('w:pBdr'))):border.getparent().remove(border)
    doc.styles['Normal'].font.size=Pt(11)
    doc.styles['Normal'].paragraph_format.line_spacing=1.12
    doc.styles['Normal'].paragraph_format.space_after=Pt(7)
    for name,size in [('Title',25),('Subtitle',14),('Heading 1',17),('Heading 2',12)]:
        doc.styles[name].font.size=Pt(size)
        doc.styles[name].paragraph_format.space_after=Pt(9)
    footer=sec.footer.paragraphs[0];footer.alignment=WD_ALIGN_PARAGRAPH.RIGHT
    field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'PAGE');footer._p.append(field)
    doc.core_properties.title=content['title'];doc.core_properties.author="PAAD research"
    def para(s,style=None):return doc.add_paragraph(s,style)
    def inputs(names):
        for i,name in enumerate(names,1):
            p0=para(f"Input {i}. {name}")
            for r in p0.runs:r.font.size=Pt(8)
    def heading(text):
        h=doc.add_heading(text,level=1);h.paragraph_format.page_break_before=True
    def picture(name):
        path=root/'figures'/name
        with Image.open(path) as im:w,h=im.size
        width=min(7.05,8.1*w/h)
        p0=doc.add_paragraph();p0.alignment=WD_ALIGN_PARAGRAPH.CENTER
        r=p0.add_run();r.add_picture(str(path),width=Inches(width))
        for node in r._r.iter(qn('wp:docPr')):node.set('descr',name.replace('_',' '))
    def result_table(t):
        heading(t['id']+'. '+t['title'])
        tbl=doc.add_table(rows=1,cols=len(t['headers']));tbl.autofit=False
        n=len(t['headers'])
        widths={3:[3.1,1.5,2.4],4:[3.0,1.0,1.5,1.5],5:[2.65,.5,1.75,1.05,1.05],6:[2.0,1.0,1.0,1.0,1.0,1.0]}[n]
        widths={
            'Table 1':[3.1,2.6,1.3],
            'Supplementary Table 2':[1.15,2.35,.7,1.4,1.4],
            'Supplementary Table 4':[1.5,1.5,1.65,.65,1.7],
            'Supplementary Table 8':[1.9,2.0,.7,2.4],
            'Supplementary Table 9':[2.55,.5,1.25,1.35,1.35],
            'Supplementary Table 10':[2.7,.75,3.55],
            'Supplementary Table 12':[3.05,.5,2.1,1.35],
        }.get(t['id'],widths)
        for j,width in enumerate(widths):tbl.columns[j].width=Inches(width)
        for j,h in enumerate(t['headers']):tbl.rows[0].cells[j].text=h
        hdr=OxmlElement('w:tblHeader');tbl.rows[0]._tr.get_or_add_trPr().append(hdr)
        for row in t['rows']:
            cells=tbl.add_row().cells
            for j,value in enumerate(row):cells[j].text=str(value)
        for ri,row in enumerate(tbl.rows):
            row._tr.get_or_add_trPr().append(OxmlElement('w:cantSplit'))
            for j,cell in enumerate(row.cells):
                cell.width=Inches(widths[j]);cell.vertical_alignment=1
                tcpr=cell._tc.get_or_add_tcPr();borders=OxmlElement('w:tcBorders')
                for edge in ['top','left','bottom','right']:
                    el=OxmlElement('w:'+edge);el.set(qn('w:val'),'single' if (ri==0 and edge in ['top','bottom']) or (ri==len(tbl.rows)-1 and edge=='bottom') else 'nil');el.set(qn('w:sz'),'6');el.set(qn('w:color'),'000000');borders.append(el)
                tcpr.append(borders)
                margins=OxmlElement('w:tcMar')
                for edge in ['top','bottom','left','right']:
                    el=OxmlElement('w:'+edge);el.set(qn('w:w'),'35');el.set(qn('w:type'),'dxa');margins.append(el)
                tcpr.append(margins)
                for pp in cell.paragraphs:
                    pp.paragraph_format.space_after=Pt(0);pp.paragraph_format.line_spacing=1.0
                    pp.alignment=WD_ALIGN_PARAGRAPH.LEFT if j==0 else WD_ALIGN_PARAGRAPH.CENTER
                    for run in pp.runs:run.font.size=Pt(9 if n>=5 else 9.5);run.bold=ri==0
        para(t['note']);inputs(t['inputs'])
    para(content['title'],'Title');para(content['subtitle'],'Subtitle')
    if content.get('review_note'):para(content['review_note'])
    para("본 보고서는 주요 유전자와 병리 소견, 전체생존의 연관성을 평가하고 기존 임상병리 평가에 유전체 정보를 추가했을 때 얻는 예후적 가치와 한계를 제시한다. DNA 복구 관련 유전자·기록된 재발·Reported TMB 및 민감도 분석은 보충자료로 구분한다.")
    para(content['discussion'][-1]['text'])
    doc.add_heading("연구 질문과 결과의 구성",1)
    for fig in content['figures']:
        if fig['id'].startswith('Figure '):para(fig['id']+'. '+fig['title'])
    heading("분석 대상과 통계 방법")
    for text in content['methods']:para(text)
    result_table(content['tables'][0])
    for fig in content['figures']:
        if fig.get('panels'):
            for panel in fig['panels']:
                heading(panel['id']+'. '+panel['title']);picture(panel['image'])
                para("그림 설명  "+panel['caption']);inputs(panel['inputs'])
        else:
            heading(fig['id']+'. '+fig['title']);picture(fig['image'])
            para("그림 설명  "+fig['method'])
        heading(fig['id']+' 설명과 해석')
        for label,key in [('분석 방법','method'),('관찰된 결과','finding'),('임상 및 생물학적 해석','interpretation'),('해석의 한계','limitation')]:
            doc.add_heading(label,2);para(fig[key])
        inputs(fig['inputs'])
        if fig['id']=='Figure 4':result_table(content['tables'][1])
    for t in content['tables'][2:]:result_table(t)
    heading("고찰과 임상적 의의")
    for part in content['discussion']:
        doc.add_heading(part['title'],2);para(part['text'])
    heading("참고문헌")
    for i,r in enumerate(content['references'],1):para(f"[{i}] {r['label']} {r['url']}")
    doc.save(target)


if __name__=='__main__':
    ap=argparse.ArgumentParser();ap.add_argument('--results',type=Path,required=True);ap.add_argument('--output',type=Path)
    ap.add_argument('--provisional',action='store_true',help='Mark as a review draft while the KRAS source decision is pending.')
    args=ap.parse_args();root=args.results.resolve()
    content=build_content(root)
    if args.provisional:
        selected='임상정보 KRAS_subtype' if content['kras_source']=='clinical' else 'MAF 기반 KRAS subtype'
        content['review_note']=f'검토본 — KRAS 출처 확정 전. {selected}을 임시 적용하였으며, 출처 선택 후 관련 결과를 확정해야 한다.'
    build=root/'private'/'document_build';build.mkdir(parents=True,exist_ok=True)
    (build/'content.json').write_text(json.dumps(content,ensure_ascii=False,indent=2),encoding='utf-8')
    target=args.output or root/'deliverables'/f"PAAD_Manuscript_Report_{date.today():%Y%m%d}.docx"
    target.parent.mkdir(parents=True,exist_ok=True)
    build_docx(content,root,target)
    print(target)
